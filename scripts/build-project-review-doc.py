from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "PROJECT_REVIEW_HANDOFF_20260908.docx"


def set_cell_shading(cell, fill):
    props = cell._tc.get_or_add_tcPr()
    shade = props.find(qn("w:shd"))
    if shade is None:
        shade = OxmlElement("w:shd")
        props.append(shade)
    shade.set(qn("w:fill"), fill)


def set_cell_border(cell, color="D9D9D9", size="6"):
    props = cell._tc.get_or_add_tcPr()
    borders = props.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        props.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_text(cell, text, *, bold=False, color="000000", size=9):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.line_spacing = 1.08
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    for index, header in enumerate(headers):
        table.columns[index].width = Inches(widths[index])
        cell = table.rows[0].cells[index]
        cell.width = Inches(widths[index])
        set_cell_shading(cell, "40545B")
        set_cell_text(cell, header, bold=True, color="FFFFFF", size=9)
        set_cell_border(cell)
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].width = Inches(widths[index])
            if row_index % 2 == 1:
                set_cell_shading(cells[index], "F2F5F4")
            set_cell_text(cells[index], str(value), size=8.6)
            set_cell_border(cells[index])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    paragraph.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    paragraph.paragraph_format.space_after = Pt(5)
    run = paragraph.add_run(text)
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.color.rgb = RGBColor(0, 0, 0)
    return paragraph


def add_body(doc, text, *, bold_lead=None):
    paragraph = doc.add_paragraph(style="Normal")
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.line_spacing = 1.18
    if bold_lead and text.startswith(bold_lead):
        lead = paragraph.add_run(bold_lead)
        lead.bold = True
        lead.font.name = "Microsoft YaHei"
        lead._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        run = paragraph.add_run(text[len(bold_lead):])
    else:
        run = paragraph.add_run(text)
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(10.5)
    return paragraph


def add_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.line_spacing = 1.08
    run = paragraph.add_run(text)
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(9.6)
    return paragraph


def configure_styles(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    for name, size in (("Heading 1", 16), ("Heading 2", 12.5), ("Heading 3", 11)):
        style = doc.styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.size = Pt(size)
        style.font.bold = True


def build():
    doc = Document()
    configure_styles(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_after = Pt(6)
    title_run = title.add_run("观象命盘 项目复盘与三术解盘闭环交接")
    title_run.font.name = "Microsoft YaHei"
    title_run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    title_run.font.size = Pt(24)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(0, 0, 0)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(16)
    subtitle_run = subtitle.add_run("2026年9月8日  Luna Max 整包交付")
    subtitle_run.font.name = "Microsoft YaHei"
    subtitle_run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    subtitle_run.font.size = Pt(11)
    subtitle_run.font.color.rgb = RGBColor(89, 105, 108)

    add_body(doc, "本文记录本轮将观象命盘从技术 RC 补到本地排盘、证据、解释、保存、历史查看和事实反馈闭环的结果。基础规则计算仍在设备本地，未接入广告、支付、云同步、真实身份供应商或生产 AI。")
    add_body(doc, "结论：R01–R10 的可在当前仓库完成部分已经实现并经过本地回归；正式公开发布仍受真机与签名、主体与商店合规、城市数据许可、传统术数专业复核和 Qwen 离线评测阻塞。")

    add_heading(doc, "一 交付结论", 1)
    add_body(doc, "本轮交付提交为 770d036729900572d06308079c5ffd3d4ec47acd，基于远端 main 3572de2 合入后推送。GitHub Actions 新 run 34201638220 与该提交绑定，真实结论为 success；该 CI 中 Web Export 与 Verify Web Export 均实际执行，不能复用旧 run 33768014467。")

    add_heading(doc, "二 F01 到 F12 完成矩阵", 1)
    matrix = [
        ("F01", "完成", "取消 saveReading 的 100 条静默截断；101 条记录完整保存并可重载回查。"),
        ("F02", "完成", "事务写入先落盘再更新内存；注入 setItem 失败时磁盘与内存均保持原值，多 key 删除和恢复支持回滚。"),
        ("F03", "完成", "解释引用只取语义候选节点；无关引用不再补数量，缺证据时显示未覆盖边界。"),
        ("F04", "完成", "八字 support 与 opposition 并存标记 conflict；只有显式 balance-validation 规则节点才标记 balanced。"),
        ("F05", "完成", "四术 payload 写入 inputFingerprint；六爻保存 timezone、seed、date、seedScope、起卦方式和 6/7/8/9 点事实。"),
        ("F06", "部分完成", "八字保留八层解释、证据链、候选取用和未知时辰 fail-closed；大运/流年深度事实引擎仍待专业规则确认。"),
        ("F07", "完成", "六爻支持快捷自动、六次交互、手工录入；保留用神、世应、旺衰、动变与反馈，不承诺具体应期。"),
        ("F08", "完成 基础", "紫微十二宫、对宫/三方、四化节点和逐宫查看可用；占星支持全量相位与日级近似边界。"),
        ("F09", "完成", "历史使用保存时 payload 和只读 ChartRenderer；反馈自动关联首个解读/依据，不静默重算历史。"),
        ("F10", "部分完成", "安全扫描、生产 audit、Web 10 路由导出和验证已收口；真机、证书、主体、许可与线上部署仍阻塞。"),
        ("F11", "方案完成", "Qwen3.5-9B 只形成脱敏离线评测方案，不接生产、不上传真实档案；权重和人工评测未执行。"),
        ("F12", "完成", "统一 npm test、类型、lint、Web Export、审计、产品闭环测试和本交接文档已完成。"),
    ]
    add_table(doc, ("编号", "状态", "真实交付内容"), matrix, (0.55, 0.95, 5.9))

    add_heading(doc, "三 三术真实操作路径", 1)
    add_heading(doc, "八字", 2)
    add_body(doc, "首页 → 命主 → 八字 → 选择性别、日界线、真太阳时 → 排出四柱 → 展开深度判断和本次计算依据 → 保存到记录 → 记录页打开保存时完整盘面、解释和事实反馈。未知时辰不补造时柱；候选取用只作为待复核入口。")
    add_heading(doc, "六爻", 2)
    add_body(doc, "首页 → 六爻 → 输入至少 4 字问题 → 选择用神方向 → 选择快捷自动、六次投掷或手工录入 → 逐条检查爻位、世应、旺衰、动静和 6/7/8/9 点 → 生成解释并保存。保存失败时点击重试保存原卦，不会重新随机。")
    add_heading(doc, "紫微", 2)
    add_body(doc, "首页 → 紫微 → 生成十二宫 → 点击宫位查看主星、对宫和三方 → 展开四化和主题解释 → 记录页查看保存时完整十二宫。命宫文案改为命宫定位，不再输出命宫落在命宫。")
    add_heading(doc, "占星与历史", 2)
    add_body(doc, "首页 → 星盘 → 生成本命盘 → 查看全部行星、宫位和相位；未知时辰只展示通过全天稳定性检查的日级落座。历史页只读保存时结果，修订规则需要用户主动运行当前规则复核并生成 Diff。")

    add_heading(doc, "四 数据迁移与回滚", 1)
    add_body(doc, "当前 Storage Schema 为 v3。future schema key 会进入 blocked 集合；所有可能覆盖、删除或清空该 key 的操作都会拒绝。设置页提供导出只读原始值，直接读取原始 AsyncStorage 字符串，不解码、不迁移、不覆盖。")
    add_body(doc, "单 key 和多 key 写入都通过 transactionalReplace 保存旧值，失败后逐 key 回滚；清除使用 transactionalRemove。只有成功落盘后才更新 React 状态和 refs。迁移失败或回滚失败时保留原始 key并报错，不强制清空，不静默重算历史。")

    add_heading(doc, "五 测试 构建与审计证据", 1)
    checks = [
        ("npm test", "通过 221/221，0 失败，0 跳过", "含 R01 storage、R02 语义引用、R03 指纹、R05 六爻手工事实、R06 紫微关系回归"),
        ("npm run typecheck", "通过", "TypeScript 无错误"),
        ("npm run lint", "通过", "Expo lint 无错误"),
        ("npm run build:web", "通过", "真实导出 10 个静态路由"),
        ("npm run verify:web", "通过", "10 routes、公共发布文件和 bundle 验证通过"),
        ("npm run security:scan", "通过", "7 个配置根目录无密钥命中"),
        ("npm run security:audit", "通过报告门禁", "0 critical / 9 high / 17 moderate / 0 low"),
        ("npm audit --omit=dev", "真实 exit 1", "0 critical / 9 high / 17 moderate / 0 low；详见生产审计报告"),
    ]
    add_table(doc, ("命令", "结果", "说明"), checks, (1.7, 1.4, 4.3))
    add_body(doc, "生产漏洞主要来自 Expo、Metro、Router、xcode 等传递依赖。本轮逐项记录了可达性和 SDK 57 兼容性，没有执行 npm audit fix --force。")

    add_heading(doc, "六 Qwen3 5 9B 离线评测边界", 1)
    add_body(doc, "设备事实为 AMD Ryzen 5 5600 与 RTX 4060 约 8 GB 显存，但仓库没有模型权重、推理服务或人工评测记录。计划采用 4-bit、本地脱敏 fixture、结构化 JSON 输出和双轮人工审核。模型只能润色已审核规则结果，不能改变排盘、证据、置信度、输入指纹或历史快照。当前状态是 evaluation-not-run。")

    add_heading(doc, "七 正式上线阻塞", 1)
    for item in [
        "真机 Web 与 iPhone 的安装、触控、键盘、VoiceOver、性能和签名证书未完成。",
        "手机号验证码、Apple、微信目前仅为本地原型入口，不得称为上线账号系统。",
        "城市数据的完整覆盖和再分发许可尚未确认，未知地点继续 fail-closed。",
        "八字大运/流年、紫微流派与四化、六爻传统规则需要专业审核和可引用来源。",
        "Qwen3.5-9B 权重、脱敏评测集和人工审查未完成；不接生产用户链路。",
        "发布主体、隐私与商店审核材料、线上部署环境仍需产品负责人提供。",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "八 文件索引", 1)
    add_body(doc, "工程交接：docs/PROJECT_REVIEW_HANDOFF_20260908.md。生产依赖审计：docs/PRODUCTION_AUDIT_20260908.md。Qwen 离线方案：docs/QWEN35_9B_OFFLINE_EVAL_PLAN_20260908.md。源代码、测试和 CI 配置以提交 770d036729900572d06308079c5ffd3d4ec47acd 为准。")

    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("观象命盘 · 本地规则 RC 交接")
    footer_run.font.name = "Microsoft YaHei"
    footer_run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = RGBColor(120, 130, 132)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
